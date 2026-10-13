from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from tests.conftest import (
    ALL_KNOWLEDGE_PERMS,
    READ_ONLY_PERMS,
    FakeIdentityClient,
    auth_header,
)

_PDF = b"""%PDF-1.1
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 51>>stream
BT /F1 12 Tf 20 120 Td (Neorvion knowledge pdf) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
trailer<</Root 1 0 R>>
%%EOF
"""


def _create_base(client: TestClient, identity: FakeIdentityClient, tenant_id: int = 11, user_id: int = 1) -> int:
    identity.grant(user_id, tenant_id, permissions=ALL_KNOWLEDGE_PERMS)
    response = client.post(
        "/api/v1/knowledge-bases",
        json={"name": f"库-{tenant_id}"},
        headers=auth_header(user_id, tenant_id),
    )
    assert response.status_code == 200, response.text
    return int(response.json()["data"]["id"])


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("产品手册", level=1)
    document.add_paragraph("这是一段用于切片的说明文字。" * 8)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_upload_parse_chunk_and_edit(
    client: TestClient,
    identity: FakeIdentityClient,
) -> None:
    base_id = _create_base(client, identity)
    headers = auth_header(1, 11)
    uploaded = client.post(
        f"/api/v1/knowledge-bases/{base_id}/documents",
        files=[
            ("files", ("a.txt", ("第一段。\n\n第二段内容用于切片测试。" * 20).encode(), "text/plain")),
            ("files", ("b.md", "# 标题\n\nMarkdown 正文。".encode(), "text/markdown")),
        ],
        headers=headers,
    )
    assert uploaded.status_code == 200, uploaded.text
    body = uploaded.json()["data"]
    assert body["succeeded"] == 2
    assert body["failed"] == 0
    doc_id = body["items"][0]["document"]["id"]
    detail = client.get(f"/api/v1/documents/{doc_id}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["data"]["status"] == "COMPLETED"
    assert detail.json()["data"]["extracted_char_count"] > 0
    assert detail.json()["data"]["chunk_unit"] == "char"
    chunks = client.get("/api/v1/chunks", params={"document_id": doc_id}, headers=headers)
    assert chunks.status_code == 200
    items = chunks.json()["data"]["items"]
    assert items
    chunk_id = items[0]["id"]
    updated = client.patch(
        f"/api/v1/chunks/{chunk_id}",
        json={"content": "人工修订后的切片", "enabled": False},
        headers=headers,
    )
    assert updated.status_code == 200
    assert updated.json()["data"]["enabled"] is False
    assert updated.json()["data"]["content"] == "人工修订后的切片"

    rechunk = client.post(
        f"/api/v1/documents/{doc_id}/chunk",
        json={"strategy": "FIXED", "chunk_size": 120, "chunk_overlap": 20},
        headers=headers,
    )
    assert rechunk.status_code == 200, rechunk.text
    after = client.get("/api/v1/chunks", params={"document_id": doc_id}, headers=headers)
    assert after.json()["data"]["items"]
    assert all(item["enabled"] is True for item in after.json()["data"]["items"])
    assert after.json()["data"]["items"][0]["content"] != "人工修订后的切片"


def test_dedup_and_partial_batch(client: TestClient, identity: FakeIdentityClient) -> None:
    base_id = _create_base(client, identity)
    headers = auth_header(1, 11)
    first = client.post(
        f"/api/v1/knowledge-bases/{base_id}/documents",
        files=[("files", ("same.txt", b"unique-body-aaa", "text/plain"))],
        headers=headers,
    )
    assert first.json()["data"]["succeeded"] == 1
    second = client.post(
        f"/api/v1/knowledge-bases/{base_id}/documents",
        files=[
            ("files", ("same.txt", b"unique-body-aaa", "text/plain")),
            ("files", ("ok.txt", b"another-file", "text/plain")),
            ("files", ("bad.exe", b"MZ", "application/octet-stream")),
        ],
        headers=headers,
    )
    assert second.status_code == 200
    data = second.json()["data"]
    assert data["succeeded"] == 1
    assert data["failed"] == 2


def test_pdf_docx_and_cross_tenant(client: TestClient, identity: FakeIdentityClient) -> None:
    base_a = _create_base(client, identity, tenant_id=11, user_id=1)
    headers_a = auth_header(1, 11)
    pdf = client.post(
        f"/api/v1/knowledge-bases/{base_a}/documents",
        files=[("files", ("guide.pdf", _PDF, "application/pdf"))],
        headers=headers_a,
    )
    assert pdf.status_code == 200, pdf.text
    assert pdf.json()["data"]["succeeded"] == 1
    doc_id = pdf.json()["data"]["items"][0]["document"]["id"]
    docx = client.post(
        f"/api/v1/knowledge-bases/{base_a}/documents",
        files=[("files", ("guide.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
        headers=headers_a,
    )
    assert docx.json()["data"]["succeeded"] == 1

    identity.grant(2, 22, permissions=ALL_KNOWLEDGE_PERMS)
    stolen = client.get(f"/api/v1/documents/{doc_id}", headers=auth_header(2, 22))
    assert stolen.status_code == 404
    download = client.get(f"/api/v1/documents/{doc_id}/file", headers=auth_header(2, 22))
    assert download.status_code == 404
    tamper = client.get(f"/api/v1/documents/{doc_id}", headers=auth_header(1, 22))
    assert tamper.status_code in {403, 404}


def test_viewer_cannot_upload(client: TestClient, identity: FakeIdentityClient) -> None:
    identity.grant(1, 11, permissions=ALL_KNOWLEDGE_PERMS)
    base = client.post(
        "/api/v1/knowledge-bases",
        json={"name": "只读库"},
        headers=auth_header(1, 11),
    ).json()["data"]["id"]
    identity.grant(3, 11, permissions=READ_ONLY_PERMS, role="VIEWER", is_owner=False)
    denied = client.post(
        f"/api/v1/knowledge-bases/{base}/documents",
        files=[("files", ("a.txt", b"nope", "text/plain"))],
        headers=auth_header(3, 11),
    )
    assert denied.status_code == 403


def test_invalid_chunk_params_and_delete(client: TestClient, identity: FakeIdentityClient) -> None:
    base_id = _create_base(client, identity)
    headers = auth_header(1, 11)
    uploaded = client.post(
        f"/api/v1/knowledge-bases/{base_id}/documents",
        files=[("files", ("keep.txt", ("可删除文档" * 30).encode(), "text/plain"))],
        headers=headers,
    )
    doc_id = uploaded.json()["data"]["items"][0]["document"]["id"]
    bad = client.post(
        f"/api/v1/documents/{doc_id}/chunk",
        json={"strategy": "FIXED", "chunk_size": 100, "chunk_overlap": 100},
        headers=headers,
    )
    assert bad.status_code == 400
    deleted = client.delete(f"/api/v1/documents/{doc_id}", headers=headers)
    assert deleted.status_code == 200
    missing = client.get(f"/api/v1/documents/{doc_id}", headers=headers)
    assert missing.status_code == 404
    overview = client.get("/api/v1/overview", headers=headers)
    assert overview.status_code == 200
    assert overview.json()["data"]["knowledge_base_count"] == 1


def test_worker_restart_marks_running_failed(client: TestClient, identity: FakeIdentityClient) -> None:
    import asyncio

    from app.core.config import settings
    from app.core.event_loop import selector_loop_factory
    from app.tasks.worker import recover_stale_tasks
    from sqlalchemy import create_engine, text

    base_id = _create_base(client, identity)
    headers = auth_header(1, 11)
    uploaded = client.post(
        f"/api/v1/knowledge-bases/{base_id}/documents",
        files=[("files", ("restart.txt", b"worker restart fixture", "text/plain"))],
        headers=headers,
    )
    doc_id = uploaded.json()["data"]["items"][0]["document"]["id"]
    engine = create_engine(settings.database_url)
    with engine.begin() as connection:
        connection.execute(
            text(
                "UPDATE document_processing_tasks SET status = 'RUNNING', finished_at = NULL "
                "WHERE document_id = :doc_id AND task_type = 'PARSE'"
            ),
            {"doc_id": doc_id},
        )
    engine.dispose()
    asyncio.run(recover_stale_tasks(), loop_factory=selector_loop_factory)
    listed = client.get("/api/v1/tasks", params={"document_id": doc_id, "task_type": "PARSE"}, headers=headers)
    assert listed.status_code == 200
    assert listed.json()["data"]["items"]
    assert listed.json()["data"]["items"][0]["status"] == "FAILED"

    retried = client.post(
        f"/api/v1/tasks/{listed.json()['data']['items'][0]['id']}/retry",
        headers=headers,
    )
    assert retried.status_code == 200, retried.text
    assert retried.json()["data"]["status"] == "SUCCESS"
    detail = client.get(f"/api/v1/documents/{doc_id}", headers=headers)
    assert detail.json()["data"]["status"] == "COMPLETED"
